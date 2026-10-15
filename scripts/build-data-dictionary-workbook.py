"""Build the data dictionary workbook covering every record kind, page, relation, file and source.

Usage: python3 scripts/build-data-dictionary-workbook.py <output.xlsx> <catalog-modules.json> <dossier-sections.json> <stats.json> [previous.xlsx]
Curated rows live in data_dictionary_curated.py; data_dictionary_sections.py generates the rest from
the exported app configuration, live statistics and the survey schema. When a previous workbook is
given, its 'Ý kiến rà soát' / 'Ghi chú của bạn' columns are carried over and summarised.
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
import data_dictionary_curated as curated
import data_dictionary_sections as extra

HEADER = PatternFill("solid", fgColor="1C426F")
FLAG = PatternFill("solid", fgColor="FFF2CC")
NOTE = PatternFill("solid", fgColor="E2F0D9")
REVIEW = ["Giữ", "Sửa", "Bỏ", "Ẩn khỏi giao diện", "Thêm vào form", "Cần thảo luận", "Đồng ý áp dụng"]
SURVEY_SCHEMA = os.path.join(os.path.dirname(__file__), "..", "src", "data", "survey-display-schema.json")


def load(path):
    with open(path, encoding="utf-8") as handle:
        return json.load(handle)


def table(wb, name, titles, widths, rows, review=False):
    ws = wb.create_sheet(name)
    ws.append(titles + (["Ý kiến rà soát", "Ghi chú của bạn"] if review else []))
    for cell, width in zip(ws[1], widths + ([16, 34] if review else [])):
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "B2"
    for row in rows:
        ws.append(list(row))
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    if review and rows:
        choice = DataValidation(type="list", formula1='"' + ",".join(REVIEW) + '"', allow_blank=True)
        ws.add_data_validation(choice)
        column = ws.cell(1, len(titles) + 1).column_letter
        choice.add(f"{column}2:{column}{len(rows) + 1}")
    return ws


def with_stats(row, stats):
    """Fill count and example of a curated row from the live statistics when it maps to one field."""
    kind = curated.GROUP_KIND.get(row[0])
    value = stats["fields"].get(f"{kind}.{row[3]}") if kind else None
    if not value:
        return row
    filled, total, example = value
    return row[:9] + (f"{filled}/{total}", row[10] or example) + row[11:]


def answered(notes, sheet, group, key, label=None):
    """The reviewer's carried note and Claude's answer for one row."""
    choice, note, _, _ = extra.note_for(notes, sheet, group, key, label)
    proposal, impact = curated.ANSWERS.get((group if sheet == "Trường dữ liệu" else "Trường chung", key), ("", ""))
    return choice, note, proposal, impact


def main(output, modules_path, sections_path, stats_path, previous=None):
    modules, sections, stats = load(modules_path), load(sections_path), load(stats_path)
    notes = extra.carried_notes(previous)
    fields = [with_stats(row, stats) for row in curated.F] + extra.field_rows(modules, stats)
    wb = Workbook()
    guide = wb.active; guide.title = "Hướng dẫn"
    for line in [
        "TỪ ĐIỂN DỮ LIỆU — toàn bộ dữ liệu app đang dùng (database thật 14/10/2026 18:15, sau lượt dọn dữ liệu cũ: 391 hồ sơ, 9 loại).",
        "'Tổng hợp ý kiến': các ghi chú bạn đã ghi ở bản trước, kèm đề xuất cách áp dụng và tác động. Chọn 'Đồng ý áp dụng' hoặc ghi lại ý kiến.",
        "Bạn mới ghi ý kiến cho Đối tác, Đơn vị và Trường chung; các nhóm Liên danh, Bài toán chuẩn, Triển khai & hợp tác, File… chưa có ý kiến.",
        "'Nhóm dữ liệu': mọi nhóm dữ liệu trên app, kể cả Liên danh, File & tài liệu, Kết quả khảo sát, Tài khoản.",
        "'Trường dữ liệu': mỗi dòng một trường hoặc một liên kết. 'Đã có dữ liệu' = số hồ sơ có giá trị / tổng. Ô vàng = nhận xét của Claude; ô xanh = phản hồi ý kiến của bạn.",
        "'Trang hồ sơ': mỗi trang hiển thị những mục nào và thao tác được gì. 'Quan hệ': các liên kết giữa các nhóm.",
        "'Bài toán chuẩn': đủ 61 bài kèm đầu mối, số đơn vị/hoạt động/đối tác. 'Liên danh': 4 liên danh. 'Câu hỏi khảo sát': bộ câu hỏi hiển thị ở Kết quả khảo sát.",
        "Ghi ý kiến ở cột 'Ý kiến rà soát' / 'Ghi chú của bạn'; lần tạo lại sau các ý kiến này được mang sang tự động. Thêm trường mới: thêm dòng cuối sheet Trường dữ liệu.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 150

    summary = []
    for row in fields:
        choice, note, proposal, impact = answered(notes, "Trường dữ liệu", row[0], row[3], row[2])
        if note or choice:
            summary.append(("Trường dữ liệu", row[0], row[2], row[3], note or choice, proposal, impact))
    for column, *_ in curated.COMMON:
        choice, note, proposal, impact = answered(notes, "Trường chung", "Trường chung", column)
        if note or choice:
            summary.append(("Trường chung", "Mọi loại hồ sơ", column, column, note or choice, proposal, impact))
    # Notes on fields that no longer exist stay listed, with what was done about them.
    listed = {(row[1], row[3]) for row in summary} | {(row[1], row[2]) for row in summary}
    for (sheet, group, _), (choice, note, label, key) in notes.items():
        group = "Mọi loại hồ sơ" if sheet == "Trường chung" else group
        if (group, key) in listed or (group, label) in listed or (note or choice) == "":
            continue
        listed.add((group, key))
        proposal, impact = curated.ANSWERS.get(("Trường chung" if sheet == "Trường chung" else group, key), ("Đã bỏ khỏi app và database.", ""))
        summary.append((sheet, group, label or key, key, note or choice, proposal, (impact + " Trường không còn.").strip()))
    table(wb, "Tổng hợp ý kiến", ["Sheet", "Nhóm", "Trường", "Mã", "Ý kiến của bạn (bản trước)", "Đề xuất áp dụng", "Tác động / cần chốt"], [14, 16, 26, 22, 40, 50, 50], summary, review=True)
    table(wb, "Nhóm dữ liệu", ["Nhóm dữ liệu", "Mã", "Vị trí trên app", "Số lượng", "Trạng thái cho phép", "Mô tả"], [24, 22, 36, 10, 44, 50], extra.kind_rows(modules, stats), review=True)

    rows = []
    for index, row in enumerate(fields, start=1):
        choice, note, proposal, impact = answered(notes, "Trường dữ liệu", row[0], row[3], row[2])
        rows.append((index, *row, f"{proposal} {impact}".strip(), choice, note))
    ws = table(wb, "Trường dữ liệu", ["STT", "Nhóm hồ sơ", "Mục trên form", "Tên trường (hiển thị)", "Mã trường", "Kiểu dữ liệu", "Giá trị cho phép", "Bắt buộc", "Hiện trên form",
               "Hiện trên trang", "Đã có dữ liệu", "Ví dụ / phân bố", "Ý nghĩa / cách dùng", "Nhận xét của Claude", "Phản hồi ý kiến của bạn"],
               [5, 16, 18, 26, 20, 16, 34, 9, 12, 22, 12, 32, 38, 40, 44], rows, review=True)
    for r, row in enumerate(rows, start=2):
        if row[13]: ws.cell(r, 14).fill = FLAG
        if row[14]: ws.cell(r, 15).fill = NOTE
        if row[2] in ("Ẩn", "Hệ thống", "Không dùng"):
            for c in range(1, 14): ws.cell(r, c).font = Font(color="6B7A8F")

    table(wb, "Trang hồ sơ", ["Trang", "Mục", "Nội dung", "Thao tác"], [20, 30, 70, 40], extra.page_rows(sections), review=True)
    table(wb, "Quan hệ", ["Từ", "Quan hệ", "Đến", "Ý nghĩa", "Số lượng hiện có", "Quy tắc", "Nhập ở đâu trên giao diện"], [20, 30, 18, 40, 24, 34, 44], curated.RELATIONS, review=True)
    names = {m["kind"]: extra.module_name(m) for m in modules} | {"partner#consortium": "Liên danh"}
    table(wb, "Quan hệ trong database", ["Từ", "Tên quan hệ (mã)", "Đến", "Số liên kết", "Nhãn gắn trên liên kết", "Có ghi chú"], [26, 20, 26, 12, 30, 12],
          [(names.get(a, a), r, names.get(b, b), n, c, k) for a, r, b, n, c, k in stats["relations"]])
    table(wb, "Vai trò trong quan hệ", ["Nhãn", "Dùng trên", "Giá trị", "Ý nghĩa"], [20, 34, 34, 70], curated.CONTEXT, review=True)

    value_rows = [(name if i == 0 else "", value, "" if count is None else count, proposal) for name, items in curated.VALUES for i, (value, count, proposal) in enumerate(items)]
    for m in modules:
        if m["kind"] not in extra.CURATED:
            counts = stats["statuses"].get(m["kind"], {})
            value_rows += [(f"Trạng thái — {extra.module_name(m)}" if i == 0 else "", s, counts.get(s, 0), "") for i, s in enumerate(m["statuses"])]
    table(wb, "Danh mục giá trị", ["Danh mục", "Giá trị", "Số hồ sơ đang dùng", "Đề xuất theo ý kiến của bạn"], [36, 58, 16, 36], value_rows, review=True)
    table(wb, "Bài toán chuẩn", ["Mã", "Tên bài toán chuẩn", "Chương trình AI", "Nhóm bài toán", "Ưu tiên", "Mức dẫn dắt", "Năm bắt đầu", "Đơn vị đầu mối", "Số đơn vị tiềm năng",
          "Số hoạt động", "Đang POC/triển khai", "Số đối tác đã làm", "Số đối tác tiềm năng"], [9, 44, 30, 34, 8, 22, 9, 14, 10, 10, 11, 10, 10], extra.standard_rows(stats), review=True)
    table(wb, "Liên danh", ["Liên danh", "Thành viên", "Số thành viên", "Số hoạt động", "Hoạt động đã chọn đối tác đầu mối", "Số file"], [32, 50, 10, 10, 16, 8], extra.consortium_rows(stats), review=True)
    table(wb, "Câu hỏi khảo sát", ["Phần", "Mã câu", "Câu hỏi", "Kiểu trả lời"], [40, 10, 80, 14], extra.survey_rows(load(SURVEY_SCHEMA)))
    common = []
    for column, label, kind, meaning in curated.COMMON:
        choice, note, proposal, impact = answered(notes, "Trường chung", "Trường chung", column)
        common.append((column, label, kind, meaning, f"{proposal} {impact}".strip(), choice, note))
    table(wb, "Trường chung", ["Mã cột", "Tên", "Kiểu", "Ý nghĩa", "Phản hồi ý kiến của bạn"], [30, 26, 24, 64, 50], common, review=True)
    table(wb, "Số liệu dashboard", ["Nhóm", "Chỉ số", "Cách app tính"], [24, 36, 90], curated.OVERVIEW, review=True)
    docs, sources = stats["documents"], stats["sources"]
    counts = {"documents": docs["documents"], "versions": docs["versions"], "links": sum(docs["links_by_kind"].values()), "audit": stats["audit_total"],
              "sources": f"{sum(sources.values())} (" + "; ".join(f"{k} {v}" for k, v in sources.items()) + ")", "source_links": stats["source_links"],
              "memberships": "; ".join(f"{k} {v}" for k, v in stats["memberships"].items()), "registry": f"{len(modules)} loại"}
    table(wb, "Bảng hệ thống", ["Bảng", "Tên", "Cột chính", "Số dòng", "Vai trò"], [20, 22, 60, 36, 60], [(n, l, c, counts[k], r) for n, l, c, k, r in curated.SYSTEM], review=True)
    idle = extra.idle_fields(fields)
    table(wb, "Trường ẩn hoặc trống", ["Nhóm hồ sơ", "Trường", "Mã trường", "Đã có dữ liệu", "Lý do liệt kê"], [26, 34, 26, 14, 40], idle, review=True)
    wb.save(output)
    print(f"{len(fields)} trường, {len(summary)} ý kiến mang sang, {len(stats['standards'])} bài toán chuẩn, {len(idle)} trường ẩn/trống; sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main(*sys.argv[1:6])
