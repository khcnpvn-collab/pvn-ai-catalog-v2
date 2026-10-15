"""Build the review workbook for survey records "Partner · Unit" that may duplicate a unit activity.

Usage: python3 scripts/build-partner-unit-review-workbook.py <records.json> <output.xlsx>
The JSON comes from plans/261014-0912-catalog-domain-redesign/data/partner-unit-records-261014.json.
Suggestions: same unit and same partner first; otherwise activities of the same unit sharing words
with the survey scope (at least two shared words). A decision is prefilled only when exactly one
same-partner activity exists and its title does not contradict a specific survey scope.
"""
import json
import re
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

DECISIONS = ["Ghép vào hoạt động", "Giữ là hợp tác chung", "Bỏ (trùng / không còn)"]
STOP = set("và các cho của trong với theo tại hỗ trợ ứng dụng xây dựng dựng giải pháp hệ thống công ty nền tảng triển khai pvn".split())
HEADER = PatternFill("solid", fgColor="1C426F")
TODO = PatternFill("solid", fgColor="FFF2CC")


def words(text):
    plain = unicodedata.normalize("NFC", (text or "").lower())
    return {w for w in re.findall(r"[\wÀ-ỹ&]+", plain) if len(w) >= 3 and w not in STOP}


def suggest(record, activities):
    code, unit, partner, status, scope, original = record
    same = [a for a in activities if a[1] == unit and a[2] == partner]
    if same:
        return same, "Cùng đơn vị và cùng đối tác"
    scored = []
    for a in activities:
        if a[1] != unit:
            continue
        score = len(words(scope) & words(a[4]))
        if score >= 2:
            scored.append((score, a))
    scored.sort(key=lambda item: -item[0])
    return [a for _, a in scored[:3]], "Cùng đơn vị, nội dung gần giống" if scored else ""


def header(ws, titles, widths):
    ws.append(titles)
    for cell, width in zip(ws[1], widths):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def main(source, output):
    data = json.load(open(source, encoding="utf-8"))
    activities = data["activities"]
    wb = Workbook()

    guide = wb.active
    guide.title = "Hướng dẫn"
    for line in [
        "Đối soát 53 bản ghi hợp tác chung \"Đối tác · Đơn vị\" từ khảo sát (14/10/2026).",
        "Mỗi dòng ở sheet \"Đối soát\" là một bản ghi. Chọn cột Quyết định:",
        "• Ghép vào hoạt động: bản ghi này chính là hoạt động ở cột \"Mã hoạt động ghép vào\" (chọn mã ở sheet \"Danh sách hoạt động\"). Đối tác, file, mốc được chuyển sang; bản ghi chung được gỡ.",
        "• Giữ là hợp tác chung: quan hệ hợp tác chung giữa đối tác và đơn vị, không ứng với một bài toán cụ thể.",
        "• Bỏ: trùng hoặc không còn giá trị; bản ghi được gỡ khỏi danh mục (vẫn giữ trong database).",
        "Nếu ghép vào hoạt động đã có đối tác khác: hệ thống tạo thêm một hoạt động riêng cho đối tác này (mỗi hoạt động 1 đối tác).",
        "Ô vàng là dòng chưa có quyết định gợi ý. Các dòng đã điền sẵn chỉ là gợi ý, bạn sửa nếu chưa đúng.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 140

    review = wb.create_sheet("Đối soát")
    header(review, ["STT", "Mã bản ghi", "Đơn vị", "Đối tác", "Trạng thái", "Tình trạng gốc", "Nội dung khảo sát", "Gợi ý (mã · tên)", "Căn cứ gợi ý", "Quyết định", "Mã hoạt động ghép vào", "Ghi chú"],
           [5, 20, 13, 24, 15, 18, 46, 70, 22, 22, 26, 30])
    for index, record in enumerate(data["general"], start=1):
        options, reason = suggest(record, activities)
        # A specific scope sharing no word with the only same-partner activity is probably another piece of work.
        contradicts = len(record[4]) > 25 and not (words(record[4]) & words(options[0][4])) if options else True
        confident = reason == "Cùng đơn vị và cùng đối tác" and len(options) == 1 and not contradicts
        decision, target = ("Ghép vào hoạt động", options[0][0]) if confident else ("", "")
        listed = "\n".join(f"{a[0]} · {a[4]} [{a[3]}]" for a in options)
        review.append([index, record[0], record[1], record[2], record[3], record[5], record[4], listed, reason, decision, target, ""])
        row = review.max_row
        for cell in review[row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if not decision:
            review.cell(row, 10).fill = TODO
            review.cell(row, 11).fill = TODO

    listing = wb.create_sheet("Danh sách hoạt động")
    header(listing, ["Mã hoạt động", "Đơn vị", "Đối tác", "Trạng thái", "Tên hoạt động", "Bài toán chuẩn", "Loại"], [40, 13, 26, 18, 90, 12, 18])
    for activity in sorted(activities, key=lambda a: (a[1], a[4])):
        listing.append(activity)
    for row in listing.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    last = len(data["general"]) + 1
    choice = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    target = DataValidation(type="list", formula1=f"='Danh sách hoạt động'!$A$2:$A${len(activities) + 1}", allow_blank=True)
    review.add_data_validation(choice)
    review.add_data_validation(target)
    choice.add(f"J2:J{last}")
    target.add(f"K2:K{last}")
    wb.save(output)
    filled = sum(1 for row in review.iter_rows(min_row=2, values_only=True) if row[9])
    print(f"{len(data['general'])} bản ghi, {filled} đã có gợi ý quyết định, {len(activities)} hoạt động để chọn")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
