"""Generated parts of the data dictionary: the record kinds, fields of the kinds outside the curated
groups, what each dossier page shows, the standard problem catalogue, consortia, survey questions,
idle fields, and the reviewer's notes carried over from the previous workbook.
Inputs: exported module definitions, dossier sections, live statistics and the survey schema.
"""
from openpyxl import load_workbook

CURATED = {"partner", "unit", "use_case", "collaboration"}
# Name used on screen when it differs from the catalogue label.
SCREEN_NAME = {"contact": "Đầu mối liên hệ", "engagement": "Mốc làm việc", "reference": "Tài liệu tham khảo / thư mục"}
# Fields offered by the compact forms; None means every visible field.
FORM_KEYS = {"contact": ["role", "position", "department", "email", "phone", "is_primary", "work_status"],
             "engagement": ["start_date", "end_date", "period_label", "location", "participants", "responsible", "conclusions", "next_notes"]}
WHERE = {"contact": "Bảng Đầu mối trên trang đối tác / đơn vị", "engagement": "Mục Mốc làm việc trên trang đối tác, đơn vị, hoạt động",
         "expert": "Trang chuyên gia", "learning_event": "Trang đào tạo & hội thảo", "reference": "Mục File & tài liệu; trang chuyên gia"}
MENU = {"use_case": "Menu Lộ trình AI", "collaboration": "Menu Lộ trình AI", "partner": "Menu Lộ trình AI", "unit": "Menu Lộ trình AI",
        "expert": "Menu Năng lực", "learning_event": "Menu Năng lực", "contact": "Trong trang đối tác / đơn vị",
        "engagement": "Trong trang đối tác / đơn vị / hoạt động", "reference": "Trong File & tài liệu, trang chuyên gia"}
TYPE = {"text": "Văn bản", "textarea": "Văn bản dài", "select": "Danh mục", "date": "Ngày", "number": "Số", "url": "Liên kết http/https", "email": "Email"}
PAGE_LABEL = {"unit": "Đơn vị", "partner": "Đối tác", "consortium": "Liên danh", "use_case": "Bài toán chuẩn", "collaboration": "Triển khai & hợp tác",
              "engagement": "Mốc làm việc", "contact": "Đầu mối liên hệ", "learning_event": "Đào tạo & hội thảo", "expert": "Chuyên gia", "reference": "Tài liệu tham khảo"}


def module_name(module):
    return SCREEN_NAME.get(module["kind"], module["label"])


def kind_rows(modules, stats):
    counts, docs, sources = stats["counts"], stats["documents"], stats["sources"]
    rows = [(module_name(m), m["kind"], MENU.get(m["kind"], ""), counts.get(m["kind"], 0), "; ".join(m["statuses"]), m["description"]) for m in modules]
    rows.insert(4, ("Liên danh", "partner (loại 'Liên danh')", "Trong trang đối tác (không hiện ở menu Đối tác)", counts["partner#consortium"], "Như đối tác", "Nhóm nhẹ các đối tác cùng tham gia; một đối tác đầu mối."))
    rows += [("File & tài liệu", "documents", "Mục File & tài liệu trên mọi trang", docs["documents"], "", "Tệp đính kèm (PDF, Office, ảnh…), có phiên bản và thư mục."),
             ("Kết quả khảo sát", "sources", "Trang đơn vị › Kết quả khảo sát (admin)", sum(sources.values()), "", "Bản gốc phiếu khảo sát và nguồn nhập; chỉ đọc."),
             ("Tài khoản & lịch sử", "memberships; audit", "Đăng nhập; không có màn hình", stats["audit_total"], "", "Quyền người dùng và lịch sử thay đổi.")]
    return rows


def field_rows(modules, stats):
    """Rows shaped like the curated ones for every kind outside the curated groups."""
    rows, fields = [], stats["fields"]
    for module in modules:
        kind = module["kind"]
        if kind in CURATED:
            continue
        name, form_keys = module_name(module), FORM_KEYS.get(kind)
        def stat(key):
            filled, total, example = fields.get(f"{kind}.{key}", [0, stats["counts"].get(kind, 0), ""])
            return f"{filled}/{total}", example
        for key, label in (("title", "Tên"), ("status", "Trạng thái")):
            filled, example = stat(key)
            rows.append((name, "Thông tin nhận diện", label, key, "Văn bản (≤250)" if key == "title" else "Danh mục",
                         "; ".join(module["statuses"]) if key == "status" else "", "Có", "Có", WHERE.get(kind, ""), filled, example, "", ""))
        for field in module["fields"]:
            filled, example = stat(field["key"])
            on_form = not field["hidden"] and (form_keys is None or field["key"] in form_keys)
            rows.append((name, field["group"] or "Nội dung", field["label"], field["key"], TYPE.get(field["type"], field["type"]),
                         "; ".join(field["options"] or field["suggestions"] or []), "Không", "Có" if on_form else "Không",
                         "Không" if field["hidden"] else WHERE.get(kind, ""), filled, example, "", "Trống hoàn toàn" if filled.startswith("0/") else ""))
    return rows


def page_rows(sections):
    """What each dossier page shows, in order, from the sections the app renders."""
    rows = []
    for kind, items in sections.items():
        page = PAGE_LABEL.get(kind, kind)
        rows.append((page, "Thông tin chính", "Các trường của hồ sơ (xem sheet Trường dữ liệu)", "Sửa trực tiếp; 'Xem đầy đủ thông tin'"))
        for section in items:
            actions = "; ".join(sorted({a.get("label") or "Thêm mới / chọn có sẵn" for a in section["actions"]})) or "Chỉ xem"
            rows.append((page, section["label"], section["description"], actions))
        rows.append((page, "File & tài liệu", "File gắn trực tiếp và file gom từ hoạt động/mốc liên quan; thư mục", "Tải lên; tạo thư mục; gỡ; ẩn file"))
        if kind == "unit":
            rows.append((page, "Kết quả khảo sát", "Nguyên văn phiếu khảo sát của đơn vị theo 10 phần A1–C6", "Chỉ xem (admin)"))
    return rows


def standard_rows(stats):
    return [tuple(row[:7]) + (row[7] or "",) + tuple(row[8:13]) for row in stats["standards"]]


def consortium_rows(stats):
    return [(c["title"], "; ".join(c["members"]), len(c["members"]), c["activities"], f"{c['activity_lead_set']}/{c['activities']}", c["files"]) for c in stats["consortia"]]


def survey_rows(schema):
    return [(section["title"], question["id"], question["label"], question["type"]) for section in schema["sections"] for question in section["questions"]]


def idle_fields(rows):
    """Fields stored but not editable on a form, or with no data yet."""
    idle = []
    for row in rows:
        group, section, label, key, form, filled = row[0], row[1], row[2], row[3], row[7], str(row[9])
        if section in ("Hệ thống", "Thông tin nhận diện", "Đề xuất mới", "Tự tính") or group in ("File & tài liệu", "Kết quả khảo sát", "Tài khoản & lịch sử") or key.startswith("quan hệ"):
            continue
        empty = filled.startswith("0/") or filled == "0"
        reason = "; ".join(reason for reason, hit in (("Không nhập được trên form", form == "Không"), ("Chưa có dữ liệu", empty)) if hit)
        if reason:
            idle.append((group, label, key, filled, reason))
    return idle


def carried_notes(path):
    """Reviewer columns of the previous workbook, keyed by sheet, group and field key or label."""
    notes = {}
    if not path:
        return notes
    for ws in load_workbook(path, read_only=True):
        header = [cell.value for cell in next(ws.iter_rows(max_row=1))]
        if "Ý kiến rà soát" not in header:
            continue
        review = header.index("Ý kiến rà soát")
        for row in ws.iter_rows(min_row=2, values_only=True):
            choice, note = (list(row[review:review + 2]) + [None, None])[:2]
            if not choice and not note:
                continue
            if ws.title == "Trường dữ liệu":
                for ident in (row[4], row[3]):
                    if ident:
                        notes[(ws.title, row[1], ident)] = (choice or "", note or "", row[3], row[4] or "")
            else:
                notes[(ws.title, ws.title, row[0])] = (choice or "", note or "", row[1], row[0])
    return notes


def note_for(notes, sheet, group, key, label=None):
    return notes.get((sheet, group, key)) or notes.get((sheet, group, label)) or ("", "", "", "")
